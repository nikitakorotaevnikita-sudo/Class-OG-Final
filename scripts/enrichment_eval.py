"""Явные промахи поиска: сколько обращений доходит до фоллбэка в каждом варианте.

Прогоняет выборку через полный пайплайн (в процессе, без сервиса) в четырёх
вариантах и пишет Excel: лист «Сводка» и лист «Детально».

    baseline  — как сейчас
    qe        — старый query expansion (ENABLE_QUERY_EXPANSION)
    append    — обогащение + исходный вопрос (QUERY_ENRICHMENT=append)
    replace   — только обогащение (QUERY_ENRICHMENT=replace)

Запуск из корня проекта (модель эмбеддингов и база — из .env):

    venv\\Scripts\\python scripts\\enrichment_eval.py [fixtures.json] [out.xlsx] [варианты через запятую]

Тексты обращений в консоль не печатаются; лог пайплайна уходит в файл рядом с xlsx.
Журнал appeals_log.jsonl прогоном не пополняется.
"""

import contextlib
import io
import json
import statistics
import sys
import time
from datetime import datetime
from pathlib import Path

sys.path.insert(0, "src")

import classifier_agent as ca
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from query_enrichment import enrichment_text

VARIANTS = {
    "baseline": {"ENABLE_QUERY_EXPANSION": False, "QUERY_ENRICHMENT": "off"},
    "qe": {"ENABLE_QUERY_EXPANSION": True, "QUERY_ENRICHMENT": "off"},
    "append": {"ENABLE_QUERY_EXPANSION": False, "QUERY_ENRICHMENT": "append"},
    "replace": {"ENABLE_QUERY_EXPANSION": False, "QUERY_ENRICHMENT": "replace"},
}


class _NoLog:
    def log(self, **_kwargs):
        return None


def run_case(agent, case, pipeline_log) -> dict:
    started = time.time()
    row = {"id": case["id"], "title": case.get("title", ""),
           "expected": case["expected_codes"]}
    try:
        with contextlib.redirect_stdout(pipeline_log):
            result = agent.classify(case["text"])
    except Exception as e:  # прогон не останавливаем
        row.update(error=f"{type(e).__name__}: {e}"[:300], sec=round(time.time() - started, 1))
        return row
    got = [q.code for q in result.questions]
    pool = {code for codes in result.retrieval_candidates.values() for code in codes}
    expected = set(case["expected_codes"])
    row.update(
        got=got,
        top1=bool(got) and got[0] in expected,
        any_hit=bool(expected & set(got)),
        in_pool=bool(expected & pool),
        trigger=result.fallback_trigger,
        fallback_used=result.full_fallback_used,
        enrichment=" | ".join(enrichment_text(e) for e in result.query_enrichments.values()),
        sec=round(time.time() - started, 1),
        error="",
    )
    return row


def summarize(name, rows) -> list:
    ok = [r for r in rows if not r.get("error")]
    secs = [r["sec"] for r in ok]
    llm = sum(r["trigger"].startswith("LLM") for r in ok)
    sim = sum(r["trigger"].startswith("max sim") for r in ok)
    return [name, len(rows), llm + sim, llm, sim,
            sum(r["in_pool"] for r in ok), sum(r["top1"] for r in ok),
            sum(r["any_hit"] for r in ok), len(rows) - len(ok),
            round(statistics.mean(secs), 1) if secs else None,
            round(statistics.median(secs), 1) if secs else None]


def write_xlsx(path, summary, details):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    head = ["Вариант", "Обращений", "Фолбэк (явный промах)", "  по сигналу LLM",
            "  по порогу сходства", "Правильный код среди кандидатов", "Top-1 верно",
            "Хоть один код верно", "Ошибок прогона", "Среднее время, с", "Медиана, с"]
    ws.append(head)
    for row in summary:
        ws.append(row)
    ws.append([])
    ws.append(["Фолбэк — поиск не нашёл тему, модель позвали по полному классификатору. "
               "Меньше — лучше."])
    ws.append([f"Порог сходства: {ca.FULL_FALLBACK_SIM_THRESHOLD}; модель эмбеддингов: "
               f"{ca.EMBEDDING_MODEL}; база: {ca.VECTOR_DB_DIR}; "
               f"дата: {datetime.now():%Y-%m-%d %H:%M}"])

    wd = wb.create_sheet("Детально")
    wd.append(["id", "Обращение (метка выборки)", "Вариант", "Правильные коды", "Выданные коды",
               "Top-1 верно", "Хоть один верно", "Правильный среди кандидатов",
               "Фолбэк", "Признак фолбэка", "Время, с", "Ошибка",
               "Что LLM дала для поиска"])
    for variant, r in details:
        wd.append([r["id"], r["title"], variant, "\n".join(r["expected"]),
                   "\n".join(r.get("got", [])),
                   _yes(r.get("top1")), _yes(r.get("any_hit")), _yes(r.get("in_pool")),
                   _yes(bool(r.get("trigger"))), r.get("trigger", ""), r["sec"],
                   r.get("error", ""), r.get("enrichment", "")])

    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for sheet in (ws, wd):
        for cell in sheet[1]:
            cell.font, cell.fill = bold, fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 16
    wd.column_dimensions["B"].width = 34
    wd.column_dimensions["D"].width = 22
    wd.column_dimensions["E"].width = 22
    wd.column_dimensions["J"].width = 30
    wd.column_dimensions["M"].width = 80
    for row in wd.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def _yes(value) -> str:
    return "" if value is None else ("да" if value else "нет")


def main():
    fixtures = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures/test_appeals_real12.json")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else
               f"eval_results/eval_enrichment_{datetime.now():%Y%m%d_%H%M}.xlsx")
    names = sys.argv[3].split(",") if len(sys.argv) > 3 else list(VARIANTS)
    cases = json.loads(fixtures.read_text(encoding="utf-8"))

    ca.get_logger = lambda: _NoLog()
    pipeline_log = io.open(out.with_suffix(".log"), "w", encoding="utf-8")
    with contextlib.redirect_stdout(pipeline_log):
        agent = ca.ClassifierAgent()

    summary, details = [], []
    for name in names:
        for attr, value in VARIANTS[name].items():
            setattr(ca, attr, value)
        rows = []
        for case in cases:
            r = run_case(agent, case, pipeline_log)
            rows.append(r)
            details.append((name, r))
            mark = "ERR" if r.get("error") else ("FB " if r["trigger"] else "   ")
            print(f"{name:8} {r['id']:7} {mark} top1={_yes(r.get('top1')):3} "
                  f"в пуле={_yes(r.get('in_pool')):3} {r['sec']:6.1f} с", flush=True)
        summary.append(summarize(name, rows))
        write_xlsx(out, summary, details)  # промежуточно — на случай обрыва
    pipeline_log.close()
    print("готово:", out)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
